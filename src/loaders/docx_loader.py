from docx import Document
from langchain_core.documents import Document as LCDocument

def load_docx(file_path: str):
    doc = Document(file_path)
    full_text = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())

    text = "\n".join(full_text)

    return [
        LCDocument(
            page_content=text,
            metadata={
                "source": file_path,
                "document_name": file_path.split("/")[-1].split("\\")[-1],
                "page": 1
            }
        )
    ]